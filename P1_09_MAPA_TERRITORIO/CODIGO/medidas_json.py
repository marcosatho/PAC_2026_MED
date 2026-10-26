# -*- coding: utf-8 -*-
"""Etapa B de la v2: umbrales y marcas de brecha por barrio, resumen por comuna, geometrías para la web y base en Excel."""
import glob, json, os, sys, warnings
import numpy as np, pandas as pd, geopandas as gpd
warnings.filterwarnings("ignore")

D = os.path.dirname(os.path.abspath(__file__))
DM = os.path.join(D, "datos_medidas")
S = pd.read_csv(os.path.join(D, "barrios_stats.csv"), dtype={"codigo": str, "comuna": str}, encoding="utf-8-sig")
S["urbano"] = S["urbano"].astype(bool)
for c in S.columns:
    if S[c].dtype.kind in "f":
        S[c] = S[c].fillna(np.nan)

# ---------- indicadores derivados ----------
S["arb_ha"] = S["arboles"] / S["ha"]
S["sanos_pct"] = np.where(S["arboles"] > 0, S["arb_sanos"] / S["arboles"] * 100, np.nan)
S["copa_pct"] = (S["evu_arboles_ha"] / S["ha"] * 100).clip(upper=100)
S["verde_pct"] = (S["evu_ha"] / S["ha"] * 100).clip(upper=100)

# ---------- umbrales (parámetros del equipo; se escriben en la página) ----------
HA_MIN = 5.0                        # barrios menores de 5 ha no se marcan (ruido de borde y áreas institucionales)
Q_ALTO, Q_BAJO = 0.75, 0.25
urb = S[S["urbano"] & (S["ha"] >= HA_MIN)]
u_lst = urb["lst_n"].dropna()
UMB = {
    "lst": float(np.quantile(u_lst, Q_ALTO)),
    "arb_ha": float(np.quantile(urb["arb_ha"].dropna(), Q_BAJO)),
    "copa": float(np.quantile(urb["copa_pct"].dropna(), Q_BAJO)),
    "mm": float(np.quantile(S[S["ha"] >= HA_MIN]["mm_alta"].dropna(), Q_ALTO)),
}
pos = S[(S["ha"] >= HA_MIN) & (S["tor_alta_ha"] > 0)]["tor_alta_ha"]
UMB["tor_ha"] = float(np.quantile(pos, Q_ALTO))
print("umbrales:", {k: round(v, 2) for k, v in UMB.items()}, "| barrios urbanos evaluados:", len(urb), "| con LST:", len(u_lst), flush=True)

ok = S["ha"] >= HA_MIN
S["f_calor"] = (S["urbano"] & ok & (S["lst_n"] >= UMB["lst"]) & (S["copa_pct"] <= UMB["copa"])).fillna(False)
# El inventario de árboles individuales (árboles por ha) se probó como criterio y explica poco la temperatura; se muestra aparte.
R_ARB, R_COPA, R_VERDE = (float(urb["lst_n"].corr(urb[c])) for c in ("arb_ha", "copa_pct", "verde_pct"))
print("correlación con la temperatura:", round(R_ARB, 2), round(R_COPA, 2), round(R_VERDE, 2), flush=True)
S["f_ladera"] = (ok & (S["mm_alta"] >= UMB["mm"]) & (S["obras"] == 0) & (S["monit"] == 0)).fillna(False)
S["f_alerta"] = (ok & (S["tor_alta_ha"] >= UMB["tor_ha"]) & (S["sat_total"] == 0)).fillna(False)
for k in ("calor", "ladera", "alerta"):
    print("barrios marcados", k, int(S["f_" + k].sum()), flush=True)

def es(x, n=1):
    return ("%.*f" % (n, x)).replace(".", ",")

crit = {
    "calor": "Barrios urbanos con temperatura de superficie sin efecto de la pendiente de %s °C o más (cuartil superior) y, a la vez, %s %% de cobertura arbórea o menos (cuartil inferior)." % (es(UMB["lst"]), es(UMB["copa"])),
    "ladera": "Barrios y veredas con %s %% o más de su área en deslizamiento alto (cuartil superior) y ninguna obra de mitigación ni instrumento de monitoreo registrado." % es(UMB["mm"], 0),
    "alerta": "Barrios y veredas con %s ha o más de avenida torrencial alta (cuartil superior de los que tienen alguna) y ninguna alarma, sensor, punto de encuentro ni ruta de evacuación." % es(UMB["tor_ha"]),
}
crit_long = {
    "calor": crit["calor"] + " La cobertura arbórea sale de los espacios verdes urbanos de 2019. Se probó también el inventario de árboles individuales (árboles por hectárea), pero explica poco la temperatura (correlación %s, frente a %s de la cobertura arbórea y %s de los espacios verdes), porque cuenta árboles en espacio público y no parques, retiros ni lotes; por eso se muestra aparte y no define la marca. Solo se evalúan barrios urbanos de %d ha o más (%d barrios)." % (es(R_ARB, 2), es(R_COPA, 2), es(R_VERDE, 2), HA_MIN, len(urb)),
    "ladera": crit["ladera"] + " Se evalúan los barrios y veredas de %d ha o más. Que no aparezca una obra no significa que otro barrio la necesite menos: el registro de obras cubre solo el programa 2016–2022." % HA_MIN,
    "alerta": crit["alerta"] + " La alerta temprana existe hoy en seis quebradas; los barrios marcados quedan fuera de ellas. Se usa el área absoluta de amenaza torrencial alta y no el porcentaje, porque el riesgo depende de cuánta quebrada hay y no de qué tan grande es el barrio. No se tiene en cuenta el alcance de una alarma: un barrio vecino a otro con alarma aparece marcado.",
}

S.to_csv(os.path.join(D, "barrios_indicadores.csv"), index=False, encoding="utf-8-sig")   # todos los indicadores y marcas, sin redondear
# ---------- resumen por comuna y corregimiento ----------
def agg(g):
    urbana = str(g.name).isdigit() and int(g.name) <= 16
    return pd.Series({
        "arboles": int(g["arboles"].sum()),
        "arb_ha": round(float(g["arboles"].sum() / g["ha"].sum()), 2) if urbana else None,
        "obras": int(g["obras"].sum()), "monit": int(g["monit"].sum()),
        "sat": int(g["sat_total"].sum()),
        "br_calor": int(g["f_calor"].sum()), "br_ladera": int(g["f_ladera"].sum()), "br_alerta": int(g["f_alerta"].sum()),
    })
terr = S.groupby("comuna").apply(agg)
terr["br_tot"] = terr["br_calor"] + terr["br_ladera"] + terr["br_alerta"]
terr_d = {k: {kk: (None if (isinstance(vv, float) and np.isnan(vv)) else (int(vv) if isinstance(vv, (np.integer,)) else vv)) for kk, vv in row.items()} for k, row in terr.to_dict("index").items()}

# ---------- barrios para la web ----------
def r(x, n=2):
    return None if pd.isna(x) else round(float(x), n)
cols_int = ["arboles", "arb_sanos", "arb_muertos", "arb_enfermos", "pot_n", "techos_n", "obras", "monit", "monit_activos", "estudios", "alarmas", "sensores", "encuentro", "rutas"]
barrios = []
for _, b in S.iterrows():
    d = {"bid": int(b["bid"]), "codigo": b["codigo"], "nombre": str(b["nombre"]).strip(), "comuna": b["comuna"], "urbano": bool(b["urbano"]), "ha": r(b["ha"], 1),
         "lst_n": r(b["lst_n"], 1), "mm_alta": r(b["mm_alta"], 1), "inu_alta": r(b["inu_alta"], 1), "tor_alta": r(b["tor_alta"], 1), "tor_alta_ha": r(b["tor_alta_ha"], 1),
         "inc_alta_ha": r(b["inc_alta_ha"], 1), "inc_media_ha": r(b["inc_media_ha"], 1),
         "arb_ha": r(b["arb_ha"], 2), "sanos_pct": r(b["sanos_pct"], 1), "copa_pct": r(b["copa_pct"], 1), "verde_pct": r(b["verde_pct"], 1), "pot_ha": r(b["pot_ha"], 2),
         "f_calor": bool(b["f_calor"]), "f_ladera": bool(b["f_ladera"]), "f_alerta": bool(b["f_alerta"])}
    for c in cols_int:
        d[c] = int(b[c]) if not pd.isna(b[c]) else 0
    barrios.append(d)

# ---------- cortes por quintiles de los barrios urbanos (para los mapas por barrio) ----------
def edges(col):
    v = urb[col].dropna()
    e = [round(float(np.quantile(v, q)), 1) for q in (0.2, 0.4, 0.6, 0.8)]
    for i in range(1, 4):
        if e[i] <= e[i - 1]:
            e[i] = round(e[i - 1] + 0.1, 1)
    return e
EDG = {"arb_edges": edges("arb_ha"), "copa_edges": edges("copa_pct"), "verde_edges": edges("verde_pct")}

# ---------- geometrías para la web (WGS84, simplificadas solo para dibujar) ----------
def cargar(n):
    fs = sorted(glob.glob(os.path.join(DM, n, "*.geojson")))
    g = pd.concat([gpd.read_file(f) for f in fs], ignore_index=True)
    g = gpd.GeoDataFrame(g, geometry="geometry", crs=9377)
    return g

def gj(gdf, props, tol=0, nd=5):
    gdf = gdf[gdf.geometry.notna() & ~gdf.geometry.is_empty].copy()
    geoms = gdf.geometry.simplify(tol, preserve_topology=True) if tol else gdf.geometry
    geoms = gpd.GeoSeries(geoms, crs=9377).to_crs(4326)
    feats = []
    from shapely.geometry import mapping
    def rnd(c):
        return [rnd(i) for i in c] if isinstance(c, (list, tuple)) else round(c, nd)
    for (_, row), g in zip(gdf.iterrows(), geoms):
        if g.is_empty:
            continue
        m = mapping(g)
        feats.append({"type": "Feature", "properties": props(row), "geometry": {"type": m["type"], "coordinates": rnd(m["coordinates"])}})
    return {"type": "FeatureCollection", "features": feats}

def cut(x, n):
    x = "" if x is None or (isinstance(x, float) and np.isnan(x)) else str(x).replace("\r", " ").replace("\n", " ").strip()
    return x if len(x) <= n else x[: n - 1].rstrip() + "…"

def fecha(ms):
    try:
        return pd.to_datetime(ms, unit="ms").strftime("%d-%m-%Y")
    except Exception:
        return ""

ob = cargar("obras")
falt = ob[ob.geometry.isna() | ob.geometry.is_empty]
if len(falt):
    ok_ = falt[falt["x_origen_nacional"].notna()]
    ob = pd.concat([ob[~(ob.geometry.isna() | ob.geometry.is_empty)], gpd.GeoDataFrame(ok_.drop(columns="geometry"), geometry=gpd.points_from_xy(ok_["x_origen_nacional"], ok_["y_origen_nacional"]), crs=9377)], ignore_index=True)
    ob = gpd.GeoDataFrame(ob, geometry="geometry", crs=9377)
inc = cargar("incendios")
inc["clase"] = inc["magnitud"].str.lower()
techos = gpd.read_file(os.path.join(D, "techos_evu.gpkg"))
GEO = {
    "incendios": gj(inc, lambda r_: {"clase": r_["clase"]}, tol=6),
    "obras": gj(ob, lambda r_: {"t": cut(r_["obra_ejecutada"], 150), "y": str(r_["anio"]), "b": cut(r_["nombre_barrio_vereda"], 40)}),
    "monit": gj(cargar("monit"), lambda r_: {"t": cut(r_["instrumento"], 30) + " · " + cut(r_["sector"], 40), "e": cut(r_["estado"], 20), "f": fecha(r_["fecha_ult_lectura"])}),
    "estudios": gj(cargar("estudios"), lambda r_: {"t": cut(r_["titulo"], 120), "y": str(r_["anio"] or "")}, tol=8),
    "alarmas": gj(cargar("alarmas"), lambda r_: {"t": cut(r_["nombre_siata"], 50), "q": cut(r_["subcuenca"], 40)}),
    "sensores": gj(cargar("sensores"), lambda r_: {"t": cut(r_["nombre_siata"], 50), "q": cut(r_["subcuenca"], 40)}),
    "encuentro": gj(cargar("encuentro"), lambda r_: {"t": cut(r_["punto_encuentro"], 110), "q": cut(r_["quebrada"], 30)}),
    "rutas": gj(cargar("rutas"), lambda r_: {"t": cut(r_["ruta_evacuacion"], 110)}, tol=3),
    "zv_pot": gj(cargar("zv_potencial"), lambda r_: {"c": cut(r_["categoria"], 30)}, tol=1),
    "techos": gj(techos.set_crs(9377, allow_override=True) if techos.crs is None else techos.to_crs(9377), lambda r_: {"c": cut(r_["categoria_msu"], 30)}, tol=1),
    "ecoparques": gj(cargar("ecoparques"), lambda r_: {"t": cut(r_["nombre"], 60), "s": cut(r_["subcategoria"], 60)}, tol=3),
    "corredores": gj(cargar("corredores"), lambda r_: {"t": cut(r_["nombre"], 60), "a": cut(r_["actuacion"], 20)}, tol=3),
}

# ---------- totales ----------
mon = cargar("monit")
urb_all = S[S["urbano"]]
tot = {
    "arboles": int(S["arboles"].sum()), "arb_sanos_pct": float(S["arb_sanos"].sum() / S["arboles"].sum() * 100),
    "arb_urb_ha": float(urb_all["arboles"].sum() / urb_all["ha"].sum()), "n_urb": int(len(urb)), "n_barr": int((S["ha"] >= HA_MIN).sum()),
    "n_calor": int(S["f_calor"].sum()), "n_ladera": int(S["f_ladera"].sum()), "n_alerta": int(S["f_alerta"].sum()),
    "evu_ha": float(S["evu_ha"].sum()), "techos_n": len(GEO["techos"]["features"]), "techos_ha": float(S["techos_ha"].sum()),
    "pot_n": len(GEO["zv_pot"]["features"]), "pot_ha": float(S["pot_ha"].sum()),
    "obras_n": len(GEO["obras"]["features"]), "obras_reg": int(len(cargar("obras"))), "monit_n": len(GEO["monit"]["features"]), "monit_act": int((mon["estado"] == "Activo").sum()),
    "estu_n": len(GEO["estudios"]["features"]), "alarm_n": len(GEO["alarmas"]["features"]), "sens_n": len(GEO["sensores"]["features"]),
    "enc_n": len(GEO["encuentro"]["features"]), "rutas_n": len(GEO["rutas"]["features"]), "eco_n": len(GEO["ecoparques"]["features"]), "cor_n": len(GEO["corredores"]["features"]),
    "inc_alta_ha": float(S["inc_alta_ha"].sum()), "inc_media_ha": float(S["inc_media_ha"].sum()), "inc_baja_ha": float(S["inc_baja_ha"].sum()),
    "inc_n": int(len(inc)), "inc_pastos": int((inc["homo"] == "Pastos enmalezados").sum()),
    "umb": UMB, "corr": {"arb": R_ARB, "copa": R_COPA, "verde": R_VERDE}, "crit": crit, "crit_long": crit_long, "hamin": HA_MIN,
}
tot.update(EDG)

# ---------- los más críticos de cada brecha ----------
def top(flag, by, asc=False, n=6):
    d = S[S[flag]].sort_values(by, ascending=asc)
    return [int(i) for i in d["bid"].head(n)]
TOP = {"calor": top("f_calor", "lst_n"), "ladera": top("f_ladera", "mm_alta"), "alerta": top("f_alerta", "tor_alta_ha")}

out = {"med": {"tot": tot, "barrios": barrios, "terr": terr_d, "top": TOP, "geo": GEO}, "bid_order": S["bid"].astype(int).tolist(), "cod_order": S["codigo"].tolist()}
p = os.path.join(D, "medidas_datos.json")
json.dump(out, open(p, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"), default=float)
print("medidas_datos.json:", round(os.path.getsize(p) / 1048576, 2), "MB", flush=True)
for k, v in GEO.items():
    print("  ", k, len(v["features"]), "| MB", round(len(json.dumps(v, separators=(",", ":"))) / 1048576, 3), flush=True)

# ---------- base por barrio en Excel (para compartir) ----------
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
terr_nom = {}
try:
    d1 = json.load(open(os.path.join(os.path.dirname(D), "mapa", "mapa_datos.json"), encoding="utf-8"))
    terr_nom = {t["codigo"]: t["nombre"] for t in d1["terr"]}
except Exception:
    pass
cols = [("codigo", "Código"), ("nombre", "Barrio o vereda"), ("comuna", "Código de comuna o corregimiento"), ("_com", "Comuna o corregimiento"), ("urbano", "Urbano"), ("ha", "Área (ha)"),
        ("lst_n", "Temperatura de superficie sin pendiente (°C, 2013–2016)"), ("mm_alta", "Deslizamientos, amenaza alta (% del área, POT 2026*)"), ("inu_alta", "Inundación, amenaza alta (% del área, POT 2026*)"),
        ("tor_alta", "Torrenciales, amenaza alta (% del área, POT 2026*)"), ("tor_alta_ha", "Torrenciales, amenaza alta (ha)"), ("inc_alta_ha", "Incendios, amenaza alta (ha)"), ("inc_media_ha", "Incendios, amenaza media (ha)"),
        ("arboles", "Árboles en el inventario"), ("arb_ha", "Árboles por hectárea"), ("sanos_pct", "Árboles sanos (%)"), ("arb_muertos", "Árboles muertos"), ("arb_enfermos", "Árboles enfermos"),
        ("copa_pct", "Cobertura arbórea (% del barrio, 2019)"), ("verde_pct", "Espacios verdes (% del barrio, 2019)"), ("pot_n", "Zonas verdes con potencial verde (n)"), ("pot_ha", "Zonas verdes con potencial verde (ha)"),
        ("techos_n", "Techos verdes marcados (n)"), ("obras", "Obras de mitigación (n)"), ("monit", "Instrumentos de monitoreo (n)"), ("estudios", "Estudios y diseños (n)"),
        ("alarmas", "Alarmas (n)"), ("sensores", "Sensores de nivel (n)"), ("encuentro", "Puntos de encuentro (n)"), ("rutas", "Rutas de evacuación (n)"),
        ("f_calor", "Brecha: calor alto y poco arbolado"), ("f_ladera", "Brecha: deslizamiento alto sin obra ni monitoreo"), ("f_alerta", "Brecha: torrencial alta sin alerta temprana")]
wb = Workbook(); ws = wb.active; ws.title = "Barrios y veredas"
ws.append([c[1] for c in cols])
for b in barrios:
    fila = []
    for k, _ in cols:
        v = terr_nom.get(b["comuna"], b["comuna"]) if k == "_com" else b[k]
        if isinstance(v, bool):
            v = "sí" if v else ""
        fila.append(v)
    ws.append(fila)
hdr = PatternFill("solid", fgColor="1E463D")
for i, c in enumerate(cols, 1):
    ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF"); ws.cell(row=1, column=i).fill = hdr
    ws.cell(row=1, column=i).alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = 14 if i > 2 else 26
ws.row_dimensions[1].height = 78
ws.freeze_panes = "C2"
w2 = wb.create_sheet("Umbrales y fuentes")
w2.append(["Tema", "Detalle"])
for t_, d_ in [
    ("Qué es esta base", "Indicadores por barrio y vereda (332 polígonos oficiales de 2014) que alimentan el mapa de la pestaña 6. Los porcentajes de amenaza salen de la intersección exacta de las capas con cada barrio; no se remuestreó ninguna capa."),
    ("Brecha: calor alto y poco arbolado", crit["calor"]),
    ("Brecha: deslizamiento alto sin obra ni monitoreo", crit["ladera"]),
    ("Brecha: torrencial alta sin alerta temprana", crit["alerta"]),
    ("Lo que significa una marca", "Que el registro no muestra la medida en ese barrio. No significa que allí falte una obra ni que otro barrio no la necesite. Los umbrales son criterio del equipo."),
    ("Amenazas (*)", "POT 2026, capas preliminares reconstruidas de láminas en PDF (Cartografía 1); no son capas oficiales."),
    ("Amenaza por incendios", "Plan municipal de gestión del riesgo; sin fecha de corte publicada."),
    ("Temperatura de superficie", "Landsat de la mañana, mediana 2013–2016, sin el efecto de la pendiente (paquete «Islas de calor» del Área Metropolitana, procesado a partir de Guzmán Echavarría, 2018). No es temperatura del aire."),
    ("Árboles", "Sistema de Árbol Urbano (servidor de mapas de la Alcaldía), consulta del 8-oct-2026; 238.021 individuos en Medellín."),
    ("Espacios verdes y cobertura arbórea", "Inventario de espacios verdes urbanos del Área Metropolitana y la Universidad Nacional (2019, contrato 1344 de 2018). Recorte exacto a cada barrio."),
    ("Obras, monitoreo, estudios, alerta temprana", "Departamento Administrativo de Gestión del Riesgo de Desastres, vía servidor de mapas de la Alcaldía (consulta del 8-oct-2026). Obras: programa 2016–2022, cobertura parcial."),
    ("Límites", "Barrios y veredas de 2014 (GeoMedellín)."),
]:
    w2.append([t_, d_])
w2.column_dimensions["A"].width = 44; w2.column_dimensions["B"].width = 130
for row in w2.iter_rows(min_row=1, max_row=w2.max_row):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for c in w2[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
xl = os.path.join(D, "Base_barrios_medidas_amenaza_PAC_v2.xlsx")
wb.save(xl)
print("excel:", xl, flush=True)
