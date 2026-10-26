# -*- coding: utf-8 -*-
"""Inventario de capas del Distrito que sirven para mostrar medidas (condiciones del presente)."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

S = "Servidor de mapas de la Alcaldía (www.medellin.gov.co/servidormapas/rest/services/"
E = "Estructura: ["  # no se usa

cols = ["Grupo", "Capa", "Qué contiene", "Geometría y n.º de registros", "Vigencia", "Dónde está (servicio / capa)",
        "Qué mostraría en el mapa", "Límites a tener presentes", "Peso para la página"]

filas = [
# ---------- verde urbano y SbN ----------
["Verde urbano y soluciones basadas en la naturaleza", "Corredores verdes (trazados posibles)",
 "Tramos viales con franja peatonal ancha propuestos por el POT 2014 para corredores verdes. Campo «actuación»: Generación 49, Mantenimiento 17, Recuperación 17.",
 "Líneas, 83", "POT 2014 (adoptado 17-dic-2014)", "vivienda_ciudad_terri/VC_Dllo_Territorial_Planificado, capa 3 (misma en ordenamiento_ter/VM_34_POT_2025, capa 1)",
 "Dónde están los corredores previstos y cuáles son de recuperación o mantenimiento (es decir, ya existen).",
 "No son los corredores construidos ni las cifras oficiales (26 corredores, 92.131 m² de jardines). Son «posibles corredores» del POT 2014: mezclar con cifras oficiales exige revisión.", "Liviana"],
["Verde urbano y soluciones basadas en la naturaleza", "Ecoparques",
 "Ecoparques administrados por Medio Ambiente: 71 de quebrada y otros cuerpos de agua, 6 de cerro, 3 de borde y 1 sin subcategoría. Campo «acción»: Mantenimiento 74, Generación 5.",
 "Polígonos, 81", "Actual (sin fecha de corte publicada en la capa)", "ambiente_dllo_sost/VC_Gestion_Ambiental, capa 2",
 "Ecoparques de quebrada como solución basada en la naturaleza junto a las manchas de inundación y avenidas torrenciales.",
 "El inventario de espacio público (Siro) registra 115 ecoparques de quebrada; las dos capas no coinciden en cantidad. Hay que escoger una y declararlo.", "Liviana"],
["Verde urbano y soluciones basadas en la naturaleza", "Red de conectividad ecológica",
 "Nodos, enlaces y fragmentos: Enlace actual 525, Enlace proyectado 105, Fragmento 58, Nodo actual 48, Nodo proyectado 40.",
 "Polígonos, 776 (más 199 de «red ecológica teórica»)", "POT 2014", "ambiente_dllo_sost/VC_Ordenamiento_Amb_Territorio, capas 1 y 5",
 "Solo lo «actual» (enlaces y nodos existentes) como infraestructura verde que conecta; lo proyectado queda fuera por la decisión de quedarnos con el presente.",
 "Es un diseño de red del POT, no una medición de conectividad.", "Liviana"],
["Verde urbano y soluciones basadas en la naturaleza", "Árbol urbano (Sistema de Árbol Urbano)",
 "Individuos con especie, estado (sano 286.587; muerto 9.663; enfermo 9.403; desconocido 25.227), intervenciones y fecha. Capa aparte de árboles patrimoniales (696).",
 "Puntos, 238.021 en Medellín (330.880 en total: la capa incluye Envigado 35.902, Bello 19.851, Copacabana 8.719 y otros municipios del Área Metropolitana)", "Actual; sistema alimentado por contratos de intervención", "ambiente_dllo_sost/VC_Gestion_Ambiental, capa 0 (también en VA_SistemaArbolUrbano_Ciudadano, capa 1)",
 "Densidad de árboles por barrio o por comuna (árboles por hectárea), proporción sana y árboles muertos o enfermos. Útil para calor y para corredores.",
 "No es un censo completo: son árboles dispersos inventariados. Hay que filtrar por municipio (campo «municipio» = MEDELLÍN). La cifra de ≈215 mil que anoté antes no coincide con la actual (238.021).", "Pesada: agregar por barrio o por celda antes de mostrar"],
["Verde urbano y soluciones basadas en la naturaleza", "Espacios verdes urbanos (EVU)",
 "Polígonos de cobertura verde: pasto 3.906 ha, árboles 2.248 ha, arbustos 145 ha, jardín 45 ha, cultivos 70 ha (suma de toda la capa). Categorías: zonas verdes privadas, retiros de quebrada, parques, separadores, antejardines, franjas de amoblamiento, ecoparques. 398 polígonos marcados con cubierta verde (techos).",
 "Polígonos, 194.916", "Inventario AMVA y Universidad Nacional, 2019 (contrato 1344 de 2018)", "ambiente_dllo_sost/VA_SistemaArbolUrbano_Base, capa 5",
 "Porcentaje de verde por barrio y por tipo (público, privado, retiro de quebrada); cuántos techos verdes hay; separar parques «grises» de «verdes».",
 "Cubre toda la jurisdicción del Área Metropolitana, no solo Medellín (en Medellín se delimitaron desde 4 m²; en los demás municipios desde 20 m²). Las hectáreas de este inventario son del conjunto.", "Pesada: agregar por barrio"],
["Verde urbano y soluciones basadas en la naturaleza", "Inventario de zonas verdes",
 "Zonas verdes del polígono prioritario urbano: asociadas a movilidad 33.795, a urbanismo 10.940, a espacio público 3.440, otras 173. Cobertura dominante: grama 25.663, árboles 11.934, jardín 4.947, tierra 3.223, arbustos 1.805, piso duro 776. Marca 1.098 con «potencial verde» y 250 con cubierta verde.",
 "Polígonos, 48.348", "Actual (sin fecha de corte publicada)", "ambiente_dllo_sost/VC_Rec_Natural_Renovable, capa 4",
 "Dónde hay zonas verdes con potencial de volverse verdes (el campo «potencial verde»), y cuáles son piso duro o tierra; apoya la idea de parques grises frente a verdes.",
 "Limitado al polígono prioritario urbano. Mucho polígono es verde residual de vías (separadores, orejas de puente).", "Pesada: agregar por barrio o filtrar potencial verde"],
["Verde urbano y soluciones basadas en la naturaleza", "Jardines (Siro)",
 "Polígonos de jardín registrados por Infraestructura Física dentro del espacio público.",
 "Polígonos, 7.377", "Actual", "vivienda_ciudad_terri/VA_Siro_jardines, capa 0 (y VC_SGVial_Espacio_Publico, capa 0)",
 "Contraste con las cifras oficiales de jardines de los corredores verdes.",
 "El campo de área suma un valor muy alto; hay que verificarlo antes de usarlo.", "Media"],
# ---------- espacio público ----------
["Espacio público", "Inventario de espacio público de esparcimiento y encuentro",
 "Parques, plazas, zonas verdes recreativas y ecoparques con categoría, subcategoría (pasivo o activo) y barrio.",
 "Polígonos, 1.507", "Actualización general más reciente: 9-oct-2025", "mapas_nacionales/VC_Infraestructura_Fisica, capa 2",
 "Espacio público por comuna y por barrio (m² por habitante si se cruza con población); parques recreativos pasivos y activos.",
 "Es la versión oficial más reciente, pero distinta en conteo del EP existente del POT 2014 (3.743) y del Siro (791 parques y 1.574 zonas verdes recreativas).", "Media"],
["Espacio público", "Espacio público existente y proyectado (POT 2014)",
 "Existente: zona verde recreacional 3.249, parque recreativo 362, ecoparque 54, plaza 50, parque cívico 24, mirador 4. Proyectado: 480 polígonos.",
 "Polígonos, 3.743 existentes y 480 proyectados", "POT 2014", "vivienda_ciudad_terri/VC_Dllo_Territorial_Planificado, capas 0 y 1",
 "Solo lo existente. Complementa el inventario de 2025.", "Viejo (2014): usar el inventario 2025 como base y esta capa solo para contraste.", "Media"],
["Espacio público", "Áreas verdes asociadas y zona verde del sistema hidrográfico",
 "Zonas verdes de secciones viales y sendero peatonal (2.679) y zonas verdes públicas del sistema hidrográfico (243).",
 "Polígonos, 2.679 y 243", "POT 2014", "ordenamiento_ter/VM_12_Subs_Espacio_Publico_Esparcimiento_Encuentro_Proy, capas 5 y 6; ambiente_dllo_sost/VC_Rec_Natural_Renovable, capa 11",
 "Verde a lo largo de las quebradas y vías.", "Se solapa con el inventario de zonas verdes y el EVU.", "Media"],
["Espacio público", "Intervenciones de espacio público en 2025 (Corredores SIF)",
 "Intervenciones de la Secretaría de Infraestructura: espacio público efectivo (195 polígonos: 181 de mantenimiento y 14 de mejoramiento, todos de 2025), andenes (17), espacio público del Centro (6) y diez corredores viales con presupuesto y número de árboles.",
 "Polígonos 195, 17 y puntos 6; líneas de 10 corredores", "2025", "vivienda_ciudad_terri/VM_Corredores_SIF, capas 0–15",
 "Qué se está interviniendo hoy y dónde, con número de árboles por intervención.",
 "Cuentan como medida en ejecución, no como cobertura total. El presupuesto proyectado suma unos $36,9 mil millones, pero habría que validar que no se repita por registro.", "Liviana"],
# ---------- drenaje y obras grises ----------
["Drenaje y obras grises", "Tramos de quebrada con tipo de tramo y material de lecho",
 "Inventario por tramo de las quebradas: tipo de tramo (códigos N 6.939, B 2.318, C 1.050), material de lecho (T 4.537, C 2.960, P 2.507, S 129…), forma, funcionalidad y si es afluente (9.031) o principal (1.278).",
 "Líneas, 10.312 tramos", "Sin fecha publicada; viene del levantamiento de quebradas de la Secretaría de Medio Ambiente", "ambiente_dllo_sost/VA_Interv_Rec_Hidrico, capa 0",
 "Porcentaje de longitud de quebradas canalizadas o cubiertas por comuna: es la mejor aproximación a «obra gris» de drenaje.",
 "Los códigos no traen diccionario en el servidor. La red hídrica de 2009 usaba «canal, cobertura, natural» para el tipo de tramo, pero no está confirmado que N, B y C signifiquen eso. Hay que pedir el diccionario a Medio Ambiente antes de pintar.", "Media (10 mil líneas)"],
["Drenaje y obras grises", "Red principal de alcantarillado",
 "Colectores (1.909 líneas) e interceptores (190) dentro de las «redes de servicios» del POT 2014; 2.099 en total de alcantarillado.",
 "Líneas, 2.099", "POT 2014", "ordenamiento_ter/VM_15_Subs_Servicios_Publicos, capa 7 (filtrar tipo_servicio = Alcantarillado)",
 "Dónde está la red de drenaje pluvial y sanitario principal frente a las manchas de inundación.",
 "Solo la red principal del POT 2014. No están las redes locales, sumideros ni pozos (están en EPM y no son públicas).", "Media"],
["Drenaje y obras grises", "Obras de mitigación de la gestión del riesgo",
 "Obras terminadas entre 2016 y 2022 (55 en 2016, 11 en 2017, 24 en 2018, 17 en 2020, 17 en 2021, 7 en 2022). Palabras en la descripción: taludes 49, senderos 44, cunetas 43, muros 20, escalas 17, gaviones 6, filtros 5, canales 4.",
 "Puntos, 131", "2016–2022", "ambiente_dllo_sost/VM_DAGRD, capa 3 (misma en VM_Gestion_Riesgo_Publico, capa 14)",
 "Dónde se ha hecho obra de mitigación de ladera, por año y por tipo.",
 "Cobertura parcial: concentradas en Popular (47) y Villa Hermosa (40), obras de presupuesto participativo. El valor de la obra se repite por contrato: no se puede sumar.", "Liviana"],
["Drenaje y obras grises", "Estudios y diseños del DAGRD",
 "Polígonos de estudios geotécnicos, hidrológicos e hidráulicos y de diseño de obras, por barrio.",
 "Polígonos, 237", "2008 en adelante", "ambiente_dllo_sost/VM_DAGRD, capa 2",
 "Dónde hay conocimiento detallado del riesgo, que es una medida de conocimiento.",
 "Muestra estudios, no obras.", "Liviana"],
["Drenaje y obras grises", "Monitoreo geotécnico",
 "Instrumentos de monitoreo en laderas: 53 activos, 14 fallados y 1 no encontrado, con fecha de la última lectura y umbral.",
 "Puntos, 68", "Última lectura registrada por instrumento", "ambiente_dllo_sost/VM_DAGRD, capa 8",
 "Dónde se vigilan laderas y cuántos instrumentos funcionan.",
 "Concentrado en Santa Elena (18), Robledo (12) y San Antonio de Prado (10).", "Liviana"],
["Drenaje y obras grises", "Hidrografía (red hídrica) y retiros",
 "Red hídrica (817 tramos de drenaje) y retiros a ríos y quebradas de 10 a 60 m (1.289 polígonos); humedales y retiro a humedales; cuencas de orden cero.",
 "Líneas 817; polígonos 1.289", "POT 2014 (hidrografía 2021 en el catálogo)", "mapas_nacionales/VC_Hidrografia; ambiente_dllo_sost/VC_Rec_Natural_Renovable, capas 3, 9 y 10",
 "Contexto de agua para todas las demás capas.", "Es base, no medida.", "Media"],
# ---------- alerta ----------
["Alerta y gestión del riesgo", "Sistema de alerta temprana: alarmas, sensores de nivel, puntos de encuentro y rutas de evacuación",
 "Alarmas (11), sensores de nivel (13), puntos de encuentro (39) y rutas de evacuación (101), asociados a quebradas de respuesta rápida: Doña María (17 puntos de encuentro), La Picacha, La Iguaná, Malpaso, Potreritos y Santa Elena.",
 "Puntos 11, 13 y 39; líneas 101", "Actual", "ambiente_dllo_sost/VC_Gestion_Riesgo, capas 6, 12, 10 y 11",
 "Qué quebradas tienen alerta y cuáles no, frente a las manchas de avenidas torrenciales del POT 2026*.",
 "Cobertura limitada a pocas quebradas; es la medida no estructural más ubicable que hay hoy.", "Liviana"],
["Alerta y gestión del riesgo", "Amenaza por incendios forestales",
 "Escenarios de amenaza: baja 381, alta 21, media 7, con uso actual y tipo de cobertura.",
 "Polígonos, 409", "Plan municipal de gestión del riesgo (sin fecha de corte publicada)", "ambiente_dllo_sost/VC_Gestion_Riesgo, capa 8",
 "Cierra parcialmente el vacío de incendios: no son eventos con coordenadas, pero sí zonas de amenaza.",
 "Es amenaza, no medida. Los eventos con ubicación siguen faltando.", "Media"],
# ---------- rural y SbN ----------
["Rural y soluciones basadas en la naturaleza", "Intervenciones forestales («Más Bosques para Medellín»)",
 "Arreglos forestales con número de árboles, sistema (101 de un tipo, 37, 15 y 7 de otros), años de plantación (2010–2025) y fechas de mantenimiento. También 89 fuentes semilleras.",
 "Polígonos 160 y puntos 89", "2010–2025", "ambiente_dllo_sost/VC_Restauracion_proteccion_ambiental, capas 1 y 0 (catálogo: «Intervenciones Forestales»)",
 "Dónde hay reforestación activa, por año de siembra, y su cercanía a quebradas y a zonas de amenaza.", "Solo lo del proyecto; no incluye la reforestación de otros actores.", "Liviana"],
["Rural y soluciones basadas en la naturaleza", "Pago por servicios ambientales y predios con exención tributaria",
 "Área elegible de pago por servicios ambientales hídricos (23 polígonos) y cobertura de predios con beneficio tributario (capas anuales 2009–2025).",
 "Polígonos, 23 y capas anuales", "2009–2025", "ambiente_dllo_sost/VM_PSA_Hidrico, capa 0; VC_Predios_catastrales_beneficio_tirbutario",
 "Dónde se conserva cobertura con incentivos, como instrumento de adaptación rural.", "Son instrumentos de política; el área elegible no es el área pagada.", "Liviana"],
# ---------- contexto ----------
["Contexto de exposición y salud", "Estrato por manzana, dengue e inseguridad alimentaria por comuna, inversión pública por comuna",
 "Estrato por manzana (32.384 polígonos); dengue por comuna (casos y tasa por año desde 2008); inseguridad alimentaria (21 polígonos, 2014–2023); inversión pública por comuna 2008 en adelante.",
 "Polígonos, 32.384; 22; 21; 21", "Varias", "mapas_nacionales/VC_Distribucion_Poblacional; salud_protec_soc/VM_Evol_Dengue y VC_Inseguridad_Alimentaria; hacienda_credi_publ/VM_Inversion_Publica",
 "Quién está expuesto o es sensible en los mismos lugares donde hay medidas; deja ver brechas de equidad.", "No son medidas. El estrato por manzana es pesado.", "Estrato: pesada; el resto liviana"],
# ---------- pospuesta ----------
["Pospuesta (se mantiene fuera por ahora)", "Coberturas terrestres 2019 y 2021",
 "Coberturas del suelo por clase (Corine adaptada). Nivel 4 en suelo rural.",
 "Polígonos, 77.769 (2019), 79.436 (2021) y 19.221 (nivel 4 rural)", "2019 y 2021", "ambiente_dllo_sost/VC_Ordenamiento_Amb_Territorio, capas 19 y 20; VM_CoberturasTerrestres2021",
 "Porcentaje de vegetación y de territorio artificializado por comuna.", "Pesada; decisión de Marcos de no implementarla todavía.", "Muy pesada"],
]

faltan = [
 ["Sistemas urbanos de drenaje sostenible", "No hay capa. En la búsqueda del catálogo geográfico no aparece ninguna con ese nombre. Lo que existe son noticias: el sistema del Parque Primavera Norte (tanques subterráneos de unos 1.200 y 500 m³ que desaguan a la quebrada Juan Bobo; en construcción), el de la cancha de Villa Niza y el del Gran Parque Medellín.",
  "Digitalizar a mano tres o cuatro puntos con la fuente de prensa, marcados como «en obra» y no como existentes; o pedir la lista a la Empresa de Desarrollo Urbano y a Medio Ambiente.", "Alcaldía de Medellín, sala de prensa; Minuto30; Telemedellín; El Colombiano"],
 ["Techos verdes, jardines de lluvia, pavimentos permeables", "Solo existen las marcas de «cubierta verde» (398 polígonos en el EVU y 250 en el inventario de zonas verdes) y el «potencial verde» (1.098). No hay inventario de jardines de lluvia ni de pavimentos permeables.",
  "Usar las marcas como aproximación y decirlo.", "Capas EVU e inventario de zonas verdes"],
 ["Tanques de retención y obras hidráulicas de control de inundación", "No hay capa. Hay noticias de mantenimiento de quebradas (7.074 m lineales en 143 puntos entre 2024 y 2025) y de obras en las quebradas Altavista, Minitas y Ana Díaz.",
  "Pedir a la Secretaría de Infraestructura Física y a la Empresa de Desarrollo Urbano el inventario de obras hidráulicas con coordenadas.", "Alcaldía de Medellín, sala de prensa"],
 ["Red de alcantarillado completa y sumideros", "Solo la red principal del POT 2014. La red completa pertenece a EPM y no es de descarga pública.", "Solicitud a EPM (lista de datos con coordenadas, solo red principal y colectores críticos).", "—"],
 ["Cinturón Verde Metropolitano y Jardín Circunvalar", "No se halló una capa descargable en el Área Metropolitana ni en el servidor de la Alcaldía. La red de conectividad ecológica del POT 2014 es lo más cercano.", "Preguntar al Área Metropolitana si publica la capa del Cinturón Verde.", "Búsqueda en la web; falta confirmar en el sitio del Área Metropolitana"],
 ["Estaciones y pluviómetros con series (SIATA)", "No hay descarga pública de series. Las alarmas y sensores de nivel del SAT ya están en el servidor de la Alcaldía.", "Solicitud formal (lista de estaciones con coordenadas y serie cruda o validada).", "—"],
 ["Temperatura del aire a 2 m", "Sin capa espacial; solo temperatura de superficie (satélite) ya en el mapa.", "Se resuelve con la solicitud al SIATA.", "—"],
 ["Código de los tramos de quebrada", "Los campos N, B, C de tipo de tramo y T, C, P de material de lecho no traen diccionario.", "Pedirlo a Medio Ambiente; sin eso no se puede decir qué porcentaje está canalizado o cubierto.", "Capa VA_Interv_Rec_Hidrico"],
]

pesos = [
 ["Liviana", "Hasta unos cientos de registros o pocos miles de líneas. Se pueden mostrar tal cual (menos de 1 MB cada una)."],
 ["Media", "Entre mil y diez mil geometrías. Se pueden mostrar simplificadas; cada una suma de 1 a 3 MB."],
 ["Pesada", "Decenas o cientos de miles de elementos. No se mandan crudas: se calcula en Python el resumen por barrio o por celda y solo se envía el resumen (unos cientos de KB)."],
 ["Muy pesada", "Coberturas del suelo: se dejan fuera, como se acordó."],
 ["Regla", "Se mantiene la malla original de cada capa. No se remuestrea ni se interpola; solo se agrega de píxel o de polígono a barrio, como en la primera fase."],
]

wb = Workbook()
ws = wb.active; ws.title = "Capas disponibles"
thin = Side(style="thin", color="CCCCCC")
hdr = PatternFill("solid", fgColor="1E463D")
def estilo(ws, ncol, widths, header_row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, max_col=ncol):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3 if ncol > 4 else 2)

ws.append(cols)
for f in filas:
    ws.append(f)
estilo(ws, len(cols), [24, 30, 58, 26, 22, 42, 46, 52, 22])

w2 = wb.create_sheet("No existe o falta")
w2.append(["Qué se buscó", "Qué se encontró", "Cómo cerrarlo", "Fuente de lo hallado"])
for f in faltan:
    w2.append(f)
estilo(w2, 4, [34, 70, 56, 36])

w3 = wb.create_sheet("Peso y reglas")
w3.append(["Peso", "Qué significa para la página"])
for f in pesos:
    w3.append(f)
estilo(w3, 2, [16, 110])

w4 = wb.create_sheet("Cómo se obtuvo")
w4.append(["Dato", "Detalle"])
for r in [
    ["Fecha de la consulta", "8-oct-2026"],
    ["Servidor", "https://www.medellin.gov.co/servidormapas/rest/services (ArcGIS Server 11.5, lectura pública). Se listaron 302 servicios y 2.608 capas, y se sondearon las candidatas (conteo de registros, campos y distribución de valores sin descargar geometrías)."],
    ["Catálogo geográfico", "https://www.medellin.gov.co/giscatalogacion (metadatos de cada capa y vigencia declarada)."],
    ["Lo que no se hizo", "No se descargó ninguna geometría ni se calcularon indicadores. Los conteos son los del servidor al momento de la consulta."],
    ["Pendientes de verificar", "Filtrar el EVU por Medellín (el árbol urbano ya se contó: 238.021); diccionario de códigos de los tramos de quebrada; área de los jardines (Siro); dependencias entre inventarios de espacio público."],
]:
    w4.append(r)
estilo(w4, 2, [26, 120])

out = r"C:\Users\marco\Downloads\Inventario_mapas_medidas_Distrito_v1.xlsx"
wb.save(out)
print("guardado:", out, "| filas:", len(filas))
