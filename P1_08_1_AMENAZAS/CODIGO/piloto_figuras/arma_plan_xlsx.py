# -*- coding: utf-8 -*-
"""Excel con el plan de ubicación de las figuras del piloto (7.1 y 8.1), la renumeración, las citas de las leyendas y los hallazgos."""
import sys, os, re, json, zipfile
from lxml import etree
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from plan_figuras import P

AQUI = os.path.dirname(os.path.abspath(__file__))
RT = r"C:\Users\marco\Downloads\PAC_P1_BORRADOR_RT.docx"
OUT = r"C:\Users\marco\Downloads\Piloto_figuras_plan_de_ubicacion.xlsx"
res = json.load(open(os.path.join(AQUI, "piloto_resultado.json"), encoding="utf-8"))
num_por_k = {r["k"]: r["num"] for r in res["registro"]}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
doc = etree.fromstring(zipfile.ZipFile(RT).read("word/document.xml"))
els = list(doc.find(W + "body"))
txt = lambda e: "".join(t.text or "" for t in e.iter(W + "t")).strip()

# ---------- hoja 1: ubicación propuesta
filas1 = []
for p in P:
    px = ""
    if p["src"] and os.path.exists(p["src"]):
        w, h = Image.open(p["src"]).size
        px = f"{w}×{h}"
    filas1.append([num_por_k.get(p["k"]) or "(sin número)", p["v2"], p["sec"], p["tipo"], p["titulo"][:140] + ("…" if len(p["titulo"]) > 140 else ""),
                   f"Después del elemento {p['anchor']}: «{txt(els[p['anchor']])[:110]}…»", p["llamada"] or "—", p["leyenda"],
                   os.path.basename(p["src"]) if p["src"] else "(sin archivo)", px, "Sí" if p["placeholder"] else "No", p["obs"], ""])
enc1 = ["N.º nuevo", "N.º en el P1 v2", "Subsección", "Tipo", "Título de la leyenda", "Ubicación (ancla en el RT)", "Llamada sugerida en el texto",
        "Origen de la leyenda", "Archivo", "Píxeles", "Pendiente", "Observación", "Su decisión (dejar / mover / quitar)"]

# ---------- hoja 2: renumeración
map_exist = res["mapa_exist"]
filas2 = []
for i, e in enumerate(els):
    if e.tag == W + "p":
        m = re.match(r"^Figura (\d+)\.(\d+)\.", txt(e))
        if m and m.group(1) in ("7", "8"):
            clave = f"{m.group(1)}.{m.group(2)}"
            filas2.append(["Leyenda existente", i, f"Figura {clave}", f"Figura {map_exist[clave]}", txt(e)[:110]])
refs = []
for i, e in enumerate(els):
    if e.tag == W + "p":
        t = txt(e)
        if re.match(r"^(Figura|Tabla)\s+\d", t):
            continue
        for m in re.finditer(r"Figuras?\s+(\d+)\.(\d+)", t):
            clave = f"{m.group(1)}.{m.group(2)}"
            if clave in map_exist:
                refs.append(["Cita en el texto", i, f"Figura {clave}", f"Figura {map_exist[clave]}", t[max(0, m.start() - 50):m.end() + 40]])
filas2 += refs
enc2 = ["Qué", "Elemento del RT", "Numeración actual", "Numeración nueva", "Contexto"]

# ---------- hoja 3: citas de las leyendas frente a la lista de referencias del RT
i_ref = max(i for i, e in enumerate(els) if e.tag == W + "p" and re.match(r"^14\. Referencias", txt(e)))  # el índice del contenido también la lista
refs_txt = [txt(e) for e in els[i_ref + 1:i_ref + 120] if e.tag == W + "p" and txt(e)]
checks = [
    ("Beck et al. (2023): Köppen-Geiger", r"Beck", ["koppen"]),
    ("Karger et al. / CHELSA v2.1", r"Karger|CHELSA", ["enso"]),
    ("NOAA/CPC (índice ONI)", r"NOAA|Climate Prediction", ["enso"]),
    ("Guzmán Echavarría (2018): temperatura de la superficie", r"Guzm[aá]n", ["lst", "lst_b"]),
    ("Secretaría de Salud de Medellín (2021): casos de dengue", r"Secretar[ií]a de Salud de Medell[ií]n.{0,14}\(2021\)", ["dengue"]),
    ("Secretaría de Salud de Medellín (s. f.): riesgo en salud", r"Secretar[ií]a de Salud de Medell[ií]n.{0,14}\(s\. f\.\)", ["salud"]),
    ("Modelo digital de terreno LiDAR 2021 (Alcaldía de Medellín)", r"LiDAR|DTM|AeroEstudios", ["elev", "pend"]),
    ("Capa de amenaza de la propuesta de revisión del POT 2026 (DAP, 2026), sin letra", r"geodatabase|capa de amenaza|Geodatabase", ["mm", "inund", "torr"]),
    ("Corantioquia, AMVA y Cornare (2018): POMCA", r"Cornare\]\. \(2018\)", ["prec", "temp", "et"]),
    ("SIATA (2019)", r"\[SIATA\]\. \(2019", ["dif", "dias", "def_acu", "def_hid", "arid"]),
    ("Universidad de Antioquia (2026a)", r"Universidad de Antioquia\. \(2026a", ["riesgo_pe", "aum", "esc"]),
    ("Universidad de Antioquia (2026b)", r"Universidad de Antioquia\. \(2026b", ["riesgo_pe", "esc", "inc"]),
    ("Universidad EAFIT (2025): estudio básico de amenaza", r"Universidad EAFIT\. \(2025", ["mm", "inund", "torr"]),
    ("Aristizábal et al. (2026)", r"Aristiz[aá]bal.{0,60}2026", ["enso_ds"]),
]
filas3 = []
for etiqueta, pat, ks in checks:
    coinc = [r for r in refs_txt if re.search(pat, r)]
    filas3.append([etiqueta, ", ".join(num_por_k.get(k) or k for k in ks), "Sí" if coinc else "NO", (coinc[0][:140] if coinc else "—"),
                   "—" if coinc else "Agregar a la lista de referencias del RT (o cambiar la cita)"])
enc3 = ["Fuente citada en la leyenda", "Figuras que la usan", "¿Está en las referencias del RT?", "Entrada encontrada", "Acción"]

# ---------- hoja 4: hallazgos
protoc = [(i, txt(e)) for i, e in enumerate(els) if e.tag == W + "p" and "protocoliz" in txt(e).lower()]
hall = [
    ["Transparencia", "El RT habla de «cartografía protocolizada» " + "; ".join(f"[{i}] {re.search('.{0,60}protocoliz.{0,50}', t, re.I).group(0)}" for i, t in protoc)
     + ". El correo del DAP dice que la información aún NO está protocolizada. Las figuras v2 ya lo declaran; falta corregir esas frases del texto."],
    ["Referencias", "Seis fuentes de las leyendas no están en las referencias del RT (Beck et al. 2023, CHELSA v2.1, NOAA/CPC, Guzmán Echavarría 2018, Secretaría de Salud de Medellín 2021 y el modelo digital de terreno LiDAR 2021), y la capa de amenaza de la propuesta del POT 2026 no tiene entrada propia (necesita su letra DAP 2026x). Las demás citas sí están. Ver la hoja «Citas de las leyendas»."],
    ["Texto que falta", "El RT no tiene texto sobre evapotranspiración (7.1), ni sobre la clasificación de Köppen (7.1), ni subsección de salud sensible al clima (8.1). Las figuras de esos temas quedan sin apoyo en el texto."],
    ["Figuras por rehacer", "La figura de proporción de suelo en amenaza por movimientos en masa por territorio (P1 v2 n.º 32) se rehace con la capa oficial v2. Las dos de salud (n.º 35 y 36) no tienen subsección."],
    ["Mapas de riesgo en 8.1", "Las figuras de riesgo por precipitación extrema y por escasez hídrica son mapas de riesgo; en el RT la clase de riesgo por comuna está en 8.4 (Tabla 8.9). Valorar pasarlas allá."],
    ["Renumeración", "El RT numera las figuras por capítulo y de corrido. Con 7 figuras nuevas en 7.1 y 16 en 8.1, las figuras 7.2 a 7.5 pasan a 7.9 a 7.12 y las 8.2 a 8.5 pasan a 8.18 a 8.21; hay que actualizar las citas del texto (hoja «Renumeración»)."],
    ["Ejes en km", "Los mapas con coordenadas planas se regeneraron con los ejes en kilómetros (marcas cada 5 km, sin rotación; 10 km en los paneles angostos), como ya estaba el panel ENOS. Cambio solo de rótulos: la proyección y la geometría no cambian. Figuras originales del SIATA y gráficos sin coordenadas quedan igual."],
    ["SIATA", "Las tres figuras de déficit hídrico y aridez vienen del estudio abierto del SIATA (2019): basta citar la fuente (no hay que pedir permiso)."],
    ["Tamaño", "El piloto pesa unos 21 MB con las 23 figuras en PNG de 2160 px; en línea convendría subir versiones más livianas."],
    ["Verificación", "No hay Word ni LibreOffice en este equipo: el documento se validó con el validador XML del skill de Word (pasó) y la ubicación se comprobó con una vista previa en HTML. Falta abrirlo en Word."],
]

wb = Workbook()
HF = PatternFill("solid", fgColor="1F3A5F")


def hoja(ws, enc, filas, anchos):
    ws.append(enc)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HF, Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = "Ubicación propuesta"
hoja(ws, enc1, filas1, [9, 9, 30, 14, 60, 60, 36, 26, 40, 11, 10, 70, 22])
for row in ws.iter_rows(min_row=2):
    if row[10].value == "Sí":
        for c in row:
            c.fill = PatternFill("solid", fgColor="FFF2CC")
hoja(wb.create_sheet("Renumeración"), enc2, filas2, [18, 14, 18, 18, 90])
hoja(wb.create_sheet("Citas de las leyendas"), enc3, filas3, [58, 30, 16, 80, 52])
ws3 = wb["Citas de las leyendas"]
for row in ws3.iter_rows(min_row=2):
    if row[2].value == "NO":
        for c in row:
            c.fill = PatternFill("solid", fgColor="F8CBAD")
hoja(wb.create_sheet("Hallazgos"), ["Tema", "Detalle"], hall, [24, 150])
wb.save(OUT)
print("guardado", OUT, "| filas ubicación:", len(filas1), "| renumeración:", len(filas2), "| citas:", len(filas3), "| faltan citas:", sum(1 for f in filas3 if f[2] == "NO"))
for f in filas3:
    print(" ", f[2], f[0])
print("protocolizada:", [i for i, _ in protoc])
