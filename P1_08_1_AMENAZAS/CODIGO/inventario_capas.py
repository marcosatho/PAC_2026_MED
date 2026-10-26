# -*- coding: utf-8 -*-
"""Solo lectura: inventario de las capas de la gdb del POT (zip YDRAY) con las evidencias de fecha declaradas
(diccionario FO-GINF-041, metadata XML de ArcGIS, campos de fecha dentro de las capas, nombre de la capa y Topologia.xlsx).
Salida: JSON intermedio (el Excel lo arma arma_inventario_xlsx.py)."""
import io, os, re, json, zipfile, warnings, collections
import xml.etree.ElementTree as ET
import openpyxl
import pandas as pd
import pyogrio
from osgeo import ogr, gdal

warnings.filterwarnings("ignore")
gdal.UseExceptions()
gdal.PushErrorHandler("CPLQuietErrorHandler")
# El zip recibido del DAP vive en Drive (08_1_AMENAZAS/05_POT2026_AMENAZAS_OFICIAL_V2/01_DATOS_ORIGINALES); la ruta se pasa por V2_ZIP.
ZIP = os.environ["V2_ZIP"]
POT = "/vsizip/" + ZIP.replace("\\", "/") + "/mj_rev_pot_2025_med.gdb"
OUT = os.environ["V2_OUTJSON"]
z = zipfile.ZipFile(ZIP)


def norm(s):
    return re.sub(r"\s+", "", str(s or "")).lower()


# ---------- 1. capas, grupos y metadata XML (GDB_Items)
ds = gdal.OpenEx(POT, gdal.OF_VECTOR, open_options=["LIST_ALL_TABLES=YES"])
items = {}
li = ds.GetLayerByName("GDB_Items")
for f in li:
    items[f.GetField("Name")] = dict(path=f.GetField("Path"), doc=f.GetField("Documentation"))


def xml_info(doc):
    out = dict(creadate="", moddate="", titulo="", resumen="", proposito="", creditos="", fecha_ref="", mddate="")
    if not doc:
        return out
    try:
        r = ET.fromstring(doc)
    except Exception:
        return out

    def t(path):
        e = r.find(path)
        return (e.text or "").strip() if e is not None and e.text else ""
    out["creadate"] = t("Esri/CreaDate")
    out["moddate"] = t("Esri/ModDate")
    out["titulo"] = t("dataIdInfo/idCitation/resTitle")
    out["resumen"] = re.sub(r"<[^>]+>", " ", t("dataIdInfo/idAbs"))
    out["resumen"] = re.sub(r"\s+", " ", out["resumen"]).strip()
    out["proposito"] = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", t("dataIdInfo/idPurp"))).strip()
    out["creditos"] = t("dataIdInfo/idCredit")
    refs = []
    for tag in ("createDate", "pubDate", "reviseDate"):
        v = t(f"dataIdInfo/idCitation/date/{tag}")
        if v:
            refs.append(f"{tag}={v}")
    out["fecha_ref"] = "; ".join(refs)
    out["mddate"] = t("mdDateSt")
    return out


capas = []
for i in range(ds.GetLayerCount()):
    l = ds.GetLayer(i)
    n = l.GetName()
    if n.startswith("GDB_") or re.match(r"^T_\d+_", n):
        continue
    d = l.GetLayerDefn()
    campos = [d.GetFieldDefn(k).GetName() for k in range(d.GetFieldCount())]
    it = items.get(n, {})
    path = (it.get("path") or "").strip("\\")
    grupo = path.split("\\")[0] if "\\" in path else "(sin grupo)"
    gt = ogr.GeometryTypeToName(l.GetGeomType()) if l.GetGeomType() else "Tabla"
    capas.append(dict(capa=n, grupo=grupo, geom=gt, n=l.GetFeatureCount(), campos=campos, xml=xml_info(it.get("doc"))))
print("capas (sin sistema ni T_):", len(capas), "| con metadata XML:", sum(1 for c in capas if c["xml"]["creadate"]))

# ---------- 2. campos de fecha dentro de las capas
patron = re.compile(r"fecha|date|anio|año|vigenc|adopc|actualiz|creaci|acuerdo|decreto", re.I)
for c in capas:
    fcol = [k for k in c["campos"] if patron.search(k)]
    c["campos_fecha"] = fcol
    c["valores_fecha"] = {}
    if not fcol or c["n"] == 0:
        continue
    try:
        df = pyogrio.read_dataframe(POT, layer=c["capa"], columns=fcol, read_geometry=False)
    except Exception as e:
        c["valores_fecha"] = {"error": str(e)[:80]}
        continue
    for k in fcol:
        s = df[k].dropna()
        if len(s) == 0:
            c["valores_fecha"][k] = dict(no_nulos=0, de=len(df))
            continue
        try:
            ss = pd.to_datetime(s, errors="coerce").dropna()
            if len(ss) >= 0.5 * len(s):
                yrs = ss.dt.year.value_counts().sort_index().to_dict()
                c["valores_fecha"][k] = dict(no_nulos=int(len(s)), de=len(df), min=str(ss.min().date()), max=str(ss.max().date()),
                                             anios={str(a): int(b) for a, b in yrs.items()})
                continue
        except Exception:
            pass
        c["valores_fecha"][k] = dict(no_nulos=int(len(s)), de=len(df), ejemplos=[str(v)[:60] for v in s.astype(str).value_counts().head(4).index])

# ---------- 3. diccionario FO-GINF-041
wb = openpyxl.load_workbook(io.BytesIO(z.read("FO_GINF041_DiccDatosGeo_mj_rev_pot_2026.xlsx")), read_only=True, data_only=True)
filas = list(wb["DiccionarioDatos"].iter_rows(values_only=True))
cab = [str(c).strip() if c else "" for c in filas[6]]
ix = {n: i for i, n in enumerate(cab) if n}


def v(f, nombre):
    i = ix.get(nombre)
    return "" if i is None or f[i] is None else re.sub(r"\s+", " ", str(f[i]).replace("\xa0", " ")).strip()


dic = collections.OrderedDict()
for f in filas[7:]:
    fc = v(f, "Nombre del feature class")
    if not fc:
        continue
    k = norm(fc)
    d = dic.setdefault(k, dict(nombre_dic=fc, dataset=v(f, "Feature dataset"), alias=v(f, "Alias FC"), geom=v(f, "Geometría / Tipo Dato"),
                              cantidad=v(f, "Cantidad de elementos"), descripcion=v(f, "Descripción"), dependencia=v(f, "Dependencia"),
                              correo=v(f, "Correo de contacto"), crs=v(f, "Sistema de coordenadas"), topologia=v(f, "Topología"),
                              reglas=v(f, "Reglas topológicas"), excepciones=v(f, "Excepciones"), fechas=[], campos=[],
                              publicable_fc=v(f, "Feature Class publicable"), clasif=collections.Counter(), campos_pub=collections.Counter()))
    fe = v(f, "Fecha de elaboración")[:10]
    if fe:
        d["fechas"].append(fe)
    cn = v(f, "Nombre del campo")
    if cn:
        d["campos"].append((cn, fe))
        d["clasif"][v(f, "Clasificación")] += 1
        d["campos_pub"][v(f, "Campo publicable")] += 1
for k, d in dic.items():
    d["clasif"] = dict(d["clasif"])
    d["campos_pub"] = dict(d["campos_pub"])
print("capas en el diccionario:", len(dic))

# ---------- 4. Topologia.xlsx
wb2 = openpyxl.load_workbook(io.BytesIO(z.read("Topologia.xlsx")), read_only=True, data_only=True)
topo = {}
for r in list(wb2.worksheets[0].iter_rows(values_only=True))[3:]:
    if r[1]:
        topo.setdefault(norm(r[1]), []).append(dict(regla=str(r[3] or "").replace("\n", " ").strip(), n=r[4], justif=str(r[5] or "")[:200], resp=str(r[6] or "")))
print("capas en Topologia.xlsx:", len(topo))

json.dump(dict(capas=capas, dic=dic, topo=topo), open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
print("guardado", OUT)
