# -*- coding: utf-8 -*-
"""Arma el borrador del capítulo de prospectiva climática (P2) en Word, con el formato del borrador del P1 (estilos, tablas, pies de figura).
No modifica el archivo del P1: lo abre como plantilla y guarda una copia nueva."""
import sys, os, re, copy, json, shutil
SCR = r"C:\Users\marco\AppData\Local\Temp\claude\G--Mi-unidad-1-Investigaci-n-Hidrometereologia-NPP-VS-PET-P-Claude\446a2666-30c7-4c61-98e1-c82458a0665f\scratchpad"
P2 = SCR + r"\p2"; FIG = P2 + r"\figs"
sys.path.insert(0, P2); sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import content1, content2, content3

RT = r"C:\Users\marco\Downloads\PAC_P1_BORRADOR_RT (2).docx"
OUT_DIR = r"C:\Users\marco\Downloads\P2_prospectiva_climatica"
OUT = os.path.join(OUT_DIR, "P2_Prospectiva_climatica_borrador_v0.1.docx")
os.makedirs(OUT_DIR, exist_ok=True)
BLOCKS = content1.B + content2.B + content3.B

# ------------------------------------------------------------------ numeración y referencias cruzadas
NUM = {}; nf = nt = 0
for b in BLOCKS:
    if b[0] == "fig": nf += 1; NUM["F:" + b[1]] = f"Figura 7.{nf}"
    if b[0] == "tbl": nt += 1; NUM["T:" + b[1]] = f"Tabla 7.{nt}"
def xref(t):
    def r(m):
        k = m.group(1)
        if k not in NUM: raise KeyError(k)
        return NUM[k]
    return re.sub(r"\{([FT]:[a-zA-Z0-9_]+)\}", r, t)

doc = Document(RT)
body = doc.element.body
ps = doc.paragraphs
cover = [copy.deepcopy(ps[i]._p) for i in range(0, 7)]          # 0 espacio, 1-4 portada, 5 salto de página, 6 «Contenido»
for ch in list(body):
    if ch.tag != qn("w:sectPr"): body.remove(ch)
for rid, rel in list(doc.part.rels.items()):
    t = rel.reltype.split("/")[-1]
    if t in ("image", "comments", "commentsIds", "commentsExtended", "commentsExtensible", "people"):
        try: doc.part.drop_rel(rid)
        except Exception as e: print("no se pudo quitar", rid, t, e)
sect = body.find(qn("w:sectPr"))

# ------------------------------------------------------------------ utilidades
def style(name): return doc.styles[name]
def add_runs(par, text, size=None, color=None, base_bold=False, base_italic=False, font=None):
    text = xref(text)
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not tok: continue
        bold, ital = base_bold, base_italic
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4: tok = tok[2:-2]; bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2: tok = tok[1:-1]; ital = True
        r = par.add_run(tok); r.bold = bold or None; r.italic = ital or None
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if font: r.font.name = font; r._r.rPr.rFonts.set(qn("w:eastAsia"), font); r._r.rPr.rFonts.set(qn("w:cs"), font)
    return par
def para(text, st="Cuerpo", **kw):
    p = doc.add_paragraph(style=style(st)); add_runs(p, text, **kw); return p
def page_break():
    p = doc.add_paragraph(); r = p.add_run(); br = OxmlElement("w:br"); br.set(qn("w:type"), "page"); r._r.append(br)
def keep_next(p): p.paragraph_format.keep_with_next = True
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); tcPr.append(sh)
def cell_borders(cell, color="BFC5CC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:color"), color); x.set(qn("w:sz"), sz); x.set(qn("w:space"), "0"); b.append(x)
    tcPr.append(b)
def cell_margins(cell, top=60, left=90, bottom=60, right=90):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for e, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:w"), str(v)); x.set(qn("w:type"), "dxa"); m.append(x)
    tcPr.append(m)
def cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn("w:tcW"))
    if tcW is None: tcW = OxmlElement("w:tcW"); tcPr.insert(0, tcW)
    tcW.set(qn("w:w"), str(w)); tcW.set(qn("w:type"), "dxa")
def set_table(tbl, widths):
    tblPr = tbl._tbl.tblPr
    for ch in list(tblPr): tblPr.remove(ch)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths))); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(widths[i]))
def fill_cell(cell, text, bold=False, size=8.5, fill=None, w=None):
    if w: cell_width(cell, w)
    cell_borders(cell)
    if fill: shade(cell, fill)
    cell_margins(cell)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.05
    add_runs(p, text, size=size, base_bold=bold)
def table(caption, headers, rows, widths, source, key=None):
    n = NUM["T:" + key]
    cp = doc.add_paragraph(style=style("Cuerpo")); keep_next(cp)
    r = cp.add_run(n + "."); r.bold = True
    add_runs(cp, " " + caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); set_table(t, widths)
    for j, h in enumerate(headers): fill_cell(t.rows[0].cells[j], h, bold=True, fill="DCE6F0", w=widths[j])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)
    for i, row in enumerate(rows, start=1):
        trPr = t.rows[i]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for j, v in enumerate(row): fill_cell(t.rows[i].cells[j], v, w=widths[j])
    sp = doc.add_paragraph(style=style("Cuerpo")); sp.paragraph_format.space_after = Pt(2); sp.paragraph_format.space_before = Pt(0)
    src = doc.add_paragraph(style=style("Cuerpo")); add_runs(src, source, size=9)
def figure(key, fname, caption, width_cm):
    n = NUM["F:" + key]
    p = doc.add_paragraph(style=style("Cuerpo")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; keep_next(p)
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(os.path.join(FIG, fname), width=Cm(width_cm))
    c = doc.add_paragraph(style=style("Normal")); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    c.paragraph_format.space_after = Pt(7); c.paragraph_format.line_spacing = 1.2
    r = c.add_run(n + "."); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string("1F2933")
    for tok_i, tok in enumerate(re.split(r"(Fuente:.*)", xref(" " + caption))):
        if not tok: continue
        rr = c.add_run(tok); rr.font.name = "Arial"; rr.font.size = Pt(10.5); rr.font.color.rgb = RGBColor.from_string("1F2933")
def box(title, text):
    t = doc.add_table(rows=1, cols=1); set_table(t, [9360]); c = t.rows[0].cells[0]
    cell_width(c, 9360); cell_borders(c, color="1F4E79", sz="6"); shade(c, "EEF4FA"); cell_margins(c, 100, 140, 100, 140)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(3); add_runs(p, title, size=9.5, base_bold=True, color="1F4E79")
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.15; p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p2, text, size=9.5)
    sp = doc.add_paragraph(style=style("Cuerpo")); sp.paragraph_format.space_after = Pt(4)

# ------------------------------------------------------------------ portada y contenido
def clone(p_el): body.insert(list(body).index(sect), copy.deepcopy(p_el))
def set_first_text(p_el, text):
    ts = p_el.findall(".//" + qn("w:t"))
    ts[0].text = text
    for t in ts[1:]: t.text = ""
    for ch in list(p_el):          # quita cambios revisados que hayan quedado
        if ch.tag in (qn("w:ins"), qn("w:del")): p_el.remove(ch)
clone(cover[0])
for el, text in ((cover[1], "Actualización del Plan de Acción Climática de Medellín"), (cover[2], "Producto 2"),
                 (cover[3], "Prospectiva climática: los escenarios que sustentan las trayectorias del PAC"), (cover[4], "Borrador estructural (no definitivo) · Versión 0.1 · 26 de octubre de 2026")):
    e = copy.deepcopy(el)
    for ch in list(e):
        if ch.tag == qn("w:ins"): e.remove(ch)
    # dejar solo la primera ejecución con texto
    runs = [r for r in e.findall(qn("w:r"))]
    for r in runs[1:]: e.remove(r)
    runs[0].find(qn("w:t")).text = text
    body.insert(list(body).index(sect), e)
box("Nota de lectura del borrador", re.sub(r"^.*?\"?$", "", "") or content1.B[0][2])
page_break()
e = copy.deepcopy(cover[6]); body.insert(list(body).index(sect), e)
toc = [b for b in BLOCKS if b[0] in ("h1", "h2")]
for b in toc:
    p = doc.add_paragraph(style=style("Cuerpo")); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(5 if b[0] == "h1" else 0)
    if b[0] == "h2": p.paragraph_format.left_indent = Cm(0.8)
    add_runs(p, b[1], size=10 if b[0] == "h2" else 10.5, base_bold=(b[0] == "h1"))
page_break()

# ------------------------------------------------------------------ cuerpo
first_h1 = True
for b in BLOCKS:
    k = b[0]
    if k == "box" and b is content1.B[0]: continue         # la nota de lectura ya se puso en la portada
    if k == "h1":
        if not first_h1: page_break()
        first_h1 = False
        para(b[1], "heading 10")
    elif k == "h2": para(b[1], "heading 20")
    elif k == "h3": para(b[1], "heading 30")
    elif k == "h4": para(b[1], "heading 40")
    elif k == "p": para(b[1], "Cuerpo")
    elif k == "fig": figure(b[1], b[2], b[3], b[4])
    elif k == "tbl": table(b[2], b[3], b[4], b[5], b[6], key=b[1])
    elif k == "box": box(b[1], b[2])
    elif k == "refs":
        for t in sorted(b[1], key=lambda s: re.sub(r"[^a-záéíóúñ ]", "", s.lower())):
            p = para(t, "Cuerpo"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0)
    else: raise ValueError(k)

# ------------------------------------------------------------------ encabezado
for s in doc.sections:
    for hf in (s.header, s.first_page_header, s.even_page_header):
        for p in hf.paragraphs:
            for r in p.runs:
                if "Producto 1" in r.text or "Versión" in r.text:
                    r.text = r.text.replace("Producto 1", "Producto 2").replace("Versión 0.11", "Borrador estructural v0.1")
            if "Producto 1" in p.text:
                full = p.text.replace("Producto 1", "Producto 2").replace("Versión 0.11", "Borrador estructural v0.1")
                for r in p.runs[1:]: r.text = ""
                p.runs[0].text = full
for tag in ("w:pPrChange", "w:rPrChange", "w:sectPrChange", "w:tblPrChange", "w:trPrChange", "w:tcPrChange", "w:ins", "w:del", "w:moveFrom", "w:moveTo"):
    for el in list(body.iter(qn(tag))):
        el.getparent().remove(el)
st = doc.settings.element
for tag in ("w:trackRevisions",):
    for el in st.findall(qn(tag)): st.remove(el)
doc.core_properties.title = "Producto 2. Prospectiva climática: escenarios que sustentan las trayectorias del PAC (borrador estructural)"
doc.core_properties.author = "Equipo PAC Medellín 2026"; doc.core_properties.comments = ""
doc.save(OUT); print("guardado", OUT, round(os.path.getsize(OUT) / 1048576, 1), "MB;", nf, "figuras;", nt, "tablas")
json.dump(NUM, open(os.path.join(P2, "numeracion.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
